from __future__ import annotations

import re
import json
import os
from bisect import bisect_right
from collections import defaultdict
from concurrent.futures import ProcessPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rapidocr import RapidOCR


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data" / "anhui-2025-undergrad"
RAW_DIR = DATA_DIR / "raw"
OUT_PATH = DATA_DIR / "anhui_2025_undergraduate_scores.xlsx"
OUT_JSON_PATH = DATA_DIR / "anhui_2025_undergraduate_scores.json"

SOURCE_URLS = {
    "历史类": "https://www.ahzsks.cn/ggl/8466.htm",
    "物理类": "https://www.ahzsks.cn/ggl/8467.htm",
}

IMAGE_GROUPS = {
    "历史类": [
        "20250724221650_437.jpg",
        "20250724221650_475.jpg",
        "20250724221650_793.jpg",
        "20250724221650_461.jpg",
        "20250724221650_322.jpg",
        "20250724221651_898.jpg",
    ],
    "物理类": [
        "20250724221823_797.jpg",
        "20250724221824_694.jpg",
        "20250724221824_194.jpg",
        "20250724221824_209.jpg",
        "20250724221824_234.jpg",
        "20250724221824_746.jpg",
        "20250724221840_108.jpg",
        "20250724221841_501.jpg",
        "20250724221841_789.jpg",
        "20250724221841_469.jpg",
        "20250724221841_819.jpg",
        "20250724221842_349.jpg",
        "20250724222113_163.jpg",
    ],
}

# Pixel boundaries measured from the official 845px-wide JPG tables.
COLUMNS = {
    "院校代码": (28, 67, 5),
    "院校名称": (68, 291, 4),
    "院校专业组": (382, 675, 4),
    "投档人数": (676, 714, 5),
    "投档最低分": (715, 761, 5),
    "最低分名次": (762, 819, 5),
}

DETAIL_HEADERS = [
    "科类",
    "页码",
    "页内行号",
    "院校代码",
    "院校名称",
    "院校专业组",
    "投档人数",
    "投档最低分",
    "最低分名次",
    "来源图片",
    "校名置信度",
    "分数置信度",
]

_OCR: RapidOCR | None = None


def find_horizontal_lines(gray: np.ndarray) -> list[int]:
    black = (gray < 90).astype(np.uint8)
    y_positions = np.where(black[:, 25:820].sum(axis=1) > 650)[0]
    if len(y_positions) == 0:
        return []

    groups: list[tuple[int, int]] = []
    start = prev = int(y_positions[0])
    for y_raw in y_positions[1:]:
        y = int(y_raw)
        if y <= prev + 1:
            prev = y
            continue
        groups.append((start, prev))
        start = prev = y
    groups.append((start, prev))
    return [(start + end) // 2 for start, end in groups]


def row_intervals(gray: np.ndarray) -> list[tuple[int, int]]:
    lines = find_horizontal_lines(gray)
    if len(lines) < 2:
        raise ValueError("Could not detect table row lines")

    # The first page of each subject has a title and header row. Continuation
    # images begin directly at y=0 with a data row.
    start_line = 1 if lines[0] > 20 else 0
    return [(lines[idx], lines[idx + 1]) for idx in range(start_line, len(lines) - 1)]


def preprocess_column(
    img: np.ndarray,
    x1: int,
    x2: int,
    y1: int,
    y2: int,
    scale: int,
) -> np.ndarray:
    crop = img[y1:y2, x1:x2]
    resized = cv2.resize(
        crop,
        ((x2 - x1) * scale, (y2 - y1) * scale),
        interpolation=cv2.INTER_CUBIC,
    )
    gray = cv2.cvtColor(resized, cv2.COLOR_BGR2GRAY)
    return cv2.threshold(gray, 180, 255, cv2.THRESH_BINARY)[1]


def clean_text(value: str) -> str:
    value = re.sub(r"\s+", "", value)
    value = value.replace("（", "(").replace("）", ")")
    value = value.replace("(", "（").replace(")", "）")
    return value.strip("，,、。. ")


def clean_number(value: str) -> str:
    value = value.replace("O", "0").replace("o", "0").replace("l", "1")
    return "".join(re.findall(r"\d+", value))


def parse_int(value: str) -> int | None:
    digits = clean_number(value)
    return int(digits) if digits else None


def ocr_column(
    ocr: RapidOCR,
    img: np.ndarray,
    intervals: list[tuple[int, int]],
    x1: int,
    x2: int,
    scale: int,
    chunk_size: int = 80,
) -> dict[int, tuple[str, float]]:
    values: dict[int, tuple[str, float]] = {}
    line_starts = [start for start, _ in intervals]

    for chunk_start in range(0, len(intervals), chunk_size):
        chunk_end = min(chunk_start + chunk_size, len(intervals))
        y1 = intervals[chunk_start][0]
        y2 = intervals[chunk_end - 1][1]
        prepared = preprocess_column(img, x1, x2, y1, y2, scale)
        result = ocr(prepared)
        if result is None or result.txts is None:
            continue

        cells: dict[int, list[tuple[float, str, float]]] = defaultdict(list)
        for box, text, score in zip(result.boxes, result.txts, result.scores):
            xs = [float(point[0]) for point in box]
            ys = [float(point[1]) for point in box]
            absolute_y = y1 + ((min(ys) + max(ys)) / 2.0 / scale)
            row_idx = bisect_right(line_starts, absolute_y) - 1
            if row_idx < chunk_start or row_idx >= chunk_end:
                continue
            cells[row_idx].append((min(xs), str(text), float(score)))

        for row_idx, parts in cells.items():
            parts.sort(key=lambda item: item[0])
            joined = "".join(part[1] for part in parts)
            avg_score = sum(part[2] for part in parts) / len(parts)
            values[row_idx] = (joined, avg_score)

    return values


def extract_image_rows(ocr: RapidOCR, subject: str, page: int, filename: str) -> list[dict]:
    path = RAW_DIR / filename
    img = cv2.imdecode(np.fromfile(path, dtype=np.uint8), cv2.IMREAD_COLOR)
    if img is None:
        raise FileNotFoundError(path)

    intervals = row_intervals(cv2.cvtColor(img, cv2.COLOR_BGR2GRAY))
    column_values = {}
    for column, (x1, x2, scale) in COLUMNS.items():
        column_values[column] = ocr_column(ocr, img, intervals, x1, x2, scale)

    rows = []
    for row_idx in range(len(intervals)):
        code_raw, _ = column_values["院校代码"].get(row_idx, ("", 0.0))
        name_raw, name_score = column_values["院校名称"].get(row_idx, ("", 0.0))
        major_raw, _ = column_values["院校专业组"].get(row_idx, ("", 0.0))
        count_raw, _ = column_values["投档人数"].get(row_idx, ("", 0.0))
        score_raw, score_conf = column_values["投档最低分"].get(row_idx, ("", 0.0))
        rank_raw, _ = column_values["最低分名次"].get(row_idx, ("", 0.0))

        rows.append(
            {
                "科类": subject,
                "页码": page,
                "页内行号": row_idx + 1,
                "院校代码": clean_number(code_raw),
                "院校名称": clean_text(name_raw),
                "院校专业组": clean_text(major_raw),
                "投档人数": parse_int(count_raw),
                "投档最低分": parse_int(score_raw),
                "最低分名次": parse_int(rank_raw),
                "来源图片": filename,
                "校名置信度": round(name_score, 4) if name_score else None,
                "分数置信度": round(score_conf, 4) if score_conf else None,
            }
        )
    return rows


def get_ocr() -> RapidOCR:
    global _OCR
    if _OCR is None:
        _OCR = RapidOCR()
    return _OCR


def extract_image_task(task: tuple[int, str, int, str]) -> tuple[int, str, int, str, list[dict]]:
    index, subject, page, filename = task
    rows = extract_image_rows(get_ocr(), subject, page, filename)
    return index, subject, page, filename, rows


def best_school_rows(rows: list[dict]) -> list[dict]:
    best: dict[tuple[str, str, str], dict] = {}
    for row in rows:
        score = row["投档最低分"]
        if score is None or not row["院校名称"]:
            continue
        key = (row["科类"], row["院校代码"], row["院校名称"])
        current = best.get(key)
        if current is None or score < current["投档最低分"]:
            best[key] = row
    return sorted(best.values(), key=lambda item: (item["科类"], item["院校名称"], item["院校代码"]))


def pivot_school_rows(best_rows: list[dict]) -> list[dict]:
    grouped: dict[str, dict] = {}
    for row in best_rows:
        name = row["院校名称"]
        item = grouped.setdefault(name, {"院校名称": name})
        prefix = "历史类" if row["科类"] == "历史类" else "物理类"
        item[f"{prefix}院校代码"] = row["院校代码"]
        item[f"{prefix}最低分"] = row["投档最低分"]
        item[f"{prefix}最低分名次"] = row["最低分名次"]
        item[f"{prefix}来源专业组"] = row["院校专业组"]
    return sorted(grouped.values(), key=lambda item: item["院校名称"])


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        max_length = 0
        letter = get_column_letter(column[0].column)
        for cell in column:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, min(len(value), 42))
        ws.column_dimensions[letter].width = max(10, max_length + 2)


def append_rows(ws, headers: list[str], rows: list[dict]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header) for header in headers])
    style_sheet(ws)


def write_workbook(detail_rows: list[dict]) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    best_rows = best_school_rows(detail_rows)
    pivot_rows = pivot_school_rows(best_rows)

    pivot_headers = [
        "院校名称",
        "历史类院校代码",
        "历史类最低分",
        "历史类最低分名次",
        "历史类来源专业组",
        "物理类院校代码",
        "物理类最低分",
        "物理类最低分名次",
        "物理类来源专业组",
    ]
    append_rows(wb.create_sheet("学校文理最低分"), pivot_headers, pivot_rows)

    best_headers = [
        "科类",
        "院校代码",
        "院校名称",
        "投档最低分",
        "最低分名次",
        "来源专业组",
        "页码",
        "页内行号",
        "来源图片",
    ]
    best_for_sheet = [
        {
            **row,
            "来源专业组": row["院校专业组"],
        }
        for row in best_rows
    ]
    append_rows(wb.create_sheet("院校最低分"), best_headers, best_for_sheet)

    for subject in ("历史类", "物理类"):
        subject_rows = [row for row in detail_rows if row["科类"] == subject]
        append_rows(wb.create_sheet(f"{subject}明细"), DETAIL_HEADERS, subject_rows)

    source_ws = wb.create_sheet("来源说明")
    source_ws.append(["项目", "内容"])
    source_ws.append(["数据来源", "安徽省教育招生考试院官网"])
    source_ws.append(["历史类页面", SOURCE_URLS["历史类"]])
    source_ws.append(["物理类页面", SOURCE_URLS["物理类"]])
    source_ws.append(["官网发布时间", "2025-07-24"])
    source_ws.append(["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    source_ws.append(["明细行数", len(detail_rows)])
    source_ws.append(["院校最低分行数", len(best_rows)])
    source_ws.append(["学校汇总行数", len(pivot_rows)])
    style_sheet(source_ws)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)


def write_json(detail_rows: list[dict]) -> None:
    best_rows = best_school_rows(detail_rows)
    pivot_rows = pivot_school_rows(best_rows)
    payload = {
        "metadata": {
            "sourceName": "安徽省教育招生考试院官网",
            "sourceUrls": SOURCE_URLS,
            "publishedAt": "2025-07-24",
            "generatedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "detailRowCount": len(detail_rows),
            "bestSchoolRowCount": len(best_rows),
            "pivotSchoolRowCount": len(pivot_rows),
        },
        "detailRows": detail_rows,
        "bestSchoolRows": best_rows,
        "pivotSchoolRows": pivot_rows,
    }
    OUT_JSON_PATH.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def main() -> None:
    tasks = []
    for subject, filenames in IMAGE_GROUPS.items():
        for page, filename in enumerate(filenames, start=1):
            tasks.append((len(tasks), subject, page, filename))

    worker_count = max(1, int(os.environ.get("ANHUI_OCR_WORKERS", "3")))
    page_rows: list[list[dict] | None] = [None] * len(tasks)

    if worker_count == 1:
        for task in tasks:
            index, subject, page, filename, rows = extract_image_task(task)
            page_rows[index] = rows
            print(f"{subject} page {page}: {filename}, rows={len(rows)}", flush=True)
    else:
        with ProcessPoolExecutor(max_workers=worker_count) as executor:
            futures = {executor.submit(extract_image_task, task): task for task in tasks}
            for future in as_completed(futures):
                index, subject, page, filename, rows = future.result()
                page_rows[index] = rows
                print(f"{subject} page {page}: {filename}, rows={len(rows)}", flush=True)

    detail_rows = [row for rows in page_rows if rows is not None for row in rows]

    write_workbook(detail_rows)
    write_json(detail_rows)
    print(f"wrote {OUT_PATH}")
    print(f"wrote {OUT_JSON_PATH}")


if __name__ == "__main__":
    main()
